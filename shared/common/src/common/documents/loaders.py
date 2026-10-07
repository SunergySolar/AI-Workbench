"""Loaders — raw bytes in, a normalised ``Document`` out.

One public entry point, ``load_document``, which sniffs the kind (detect.py)
and dispatches to the per-kind loader:

    image → ``_load_image``  one page, decoded through a caller-supplied
                             decoder (so a service can keep its own
                             EXIF-aware decode path), no text layer.
    pdf   → ``_load_pdf``    per-page native text via PyMuPDF plus a raster
                             render at ``render_dpi``, capped at ``max_pages``.
    svg   → ``_load_svg``    one page, rendered by MuPDF like a one-page PDF
                             (native text from ``<text>``, a raster render
                             capped at ``max_render_pixels``), plus a warning
                             per external reference it did not draw.
    txt   → ``_load_txt``    one page, text only, no image.
    docx  → ``_load_docx``   one page, paragraphs and table cells in document
                             order, no image.

Design notes worth knowing before editing:

  * Only PDFs are multi-page. An SVG is one drawing — one page. .docx has
    no page concept until it is laid out (python-docx reads the XML, not a
    renderer), and .txt has none at all, so both collapse to a single page.
    Callers that need per-page text for those kinds have to render them
    first — out of scope here.
  * An SVG is rendered by MuPDF, which fetches NOTHING an SVG points at
    (http, ``file://``, paths, ``<use>``, CSS imports, XXE — see ``svg.py``).
    That is a property this module relies on and the tests pin; a service
    that wants remote images inlines them as ``data:`` URIs BEFORE loading.
  * Nothing in this module imports OpenCV. Images are handled with PIL and
    numpy so the package stays usable in services that only depend on one of
    the two OpenCV wheels (or neither).
  * The image decoder is injected rather than fixed. ``classifier`` passes its
    EXIF-correcting decoder so a phone photo loads the right way up; anything
    else gets ``default_image_decoder``.
  * Every PyMuPDF call holds ``_PYMUPDF_LOCK``. PyMuPDF documents itself as
    not thread-safe (one global MuPDF context), and callers run these
    functions from worker threads (``asyncio.to_thread``) — several jobs at
    once. The PDF render holds it per page, not per document, so a 20-page
    render does not stall a short ``pdf_text_regions`` lookup for seconds.

Requires the ``documents`` extra: ``pymupdf`` (PDF, SVG), ``python-docx`` (DOCX),
``pillow`` + ``numpy`` (images). Each is imported lazily inside its loader so
a consumer that only handles .txt pays for none of them.

Process flow position: called by services on upload; the ``Document`` it
returns is then optionally OCR'd (ocr.py) and searched (textmatch.py).
"""

from __future__ import annotations

import math
import threading
from typing import TYPE_CHECKING, Callable, Optional

from .detect import DocumentKind, UnsupportedDocumentError, detect_kind, guess_content_type
from .model import Document, Page
from .svg import find_external_refs

if TYPE_CHECKING:  # pragma: no cover - typing only
    import numpy as np

# raw bytes → BGR numpy array. Injected into load_document so each service can
# keep its own decode semantics (EXIF rotation, colour handling, fallbacks).
ImageDecoder = Callable[[bytes], "np.ndarray"]

# Defaults, deliberately conservative:
#   20 pages keeps a 200-page contract from turning into 200 OCR passes.
#   150 dpi is the lowest density at which 8-10pt body text survives OCR.
#   25 megapixels bounds ONE SVG render (≈75 MB as RGB). An SVG declares its
#   own size, and ``width="1000000"`` is three characters more than
#   ``width="1000"``: past the cap the zoom is scaled down so the render fits,
#   instead of MuPDF refusing (or allocating) a terapixel. PDFs are not capped
#   — large-format plan sets render at their DPI as they always have.
DEFAULT_MAX_PAGES = 20
DEFAULT_RENDER_DPI = 150
DEFAULT_MAX_RENDER_PIXELS = 25_000_000

# How many external SVG references get a warning line of their own before the
# rest are summarised as "+N more". A malicious file can carry thousands.
_MAX_REF_WARNINGS = 10
# A reported URL is cut to this many characters — a data-exfiltration URL can
# be kilobytes long, and the warning only has to identify it.
_MAX_WARNING_URL = 200

# PyMuPDF is not thread-safe; every call into it in this module holds this
# (see the module docstring). Re-entrant so a helper can call another.
_PYMUPDF_LOCK = threading.RLock()


def default_image_decoder(raw: bytes) -> "np.ndarray":
    """Decode image bytes to a BGR numpy array using PIL.

    Applies EXIF orientation so a portrait phone photo does not arrive
    sideways, converts to RGB (dropping alpha / palette), then reverses the
    channel order to the BGR convention every OpenCV consumer expects.

    Args:
        raw: JPEG or PNG bytes.

    Returns:
        BGR numpy array (H×W×3, uint8).

    Raises:
        UnsupportedDocumentError: If PIL cannot decode the bytes.
    """
    import io

    import numpy as np
    from PIL import Image, ImageOps

    try:
        pil = Image.open(io.BytesIO(raw))
        pil = ImageOps.exif_transpose(pil)
        rgb = np.array(pil.convert("RGB"))
    except Exception as exc:  # PIL raises a zoo of exception types
        raise UnsupportedDocumentError(f"Could not decode image: {exc}") from exc
    # RGB → BGR without OpenCV; .copy() because downstream code writes in place.
    return rgb[:, :, ::-1].copy()


# ---------------------------------------------------------------------------
# Per-kind loaders
# ---------------------------------------------------------------------------


def _load_image(raw: bytes, decoder: ImageDecoder) -> list[Page]:
    """One page holding the decoded image and no text layer."""
    bgr = decoder(raw)
    h, w = bgr.shape[:2]
    return [Page(index=0, image_bgr=bgr, text="", text_source="none", width=w, height=h)]


def _load_pdf(raw: bytes, max_pages: int, render_dpi: int) -> tuple[list[Page], int]:
    """Extract native text and a raster render for each PDF page.

    Both outputs matter: the text layer is exact when the PDF was generated
    digitally, and the render is what CV detectors and vision models see (and
    what OCR reads when the text layer is empty — a scan).

    Args:
        raw:        PDF bytes.
        max_pages:  Hard cap on pages loaded; the rest are counted as truncated.
        render_dpi: Raster density for the page render.

    Returns:
        ``(pages, truncated_pages)``.

    Raises:
        UnsupportedDocumentError: If PyMuPDF cannot open the stream.
    """
    import numpy as np
    import pymupdf  # PyMuPDF ≥1.24 exposes both `pymupdf` and the legacy `fitz`

    try:
        with _PYMUPDF_LOCK:
            doc = pymupdf.open(stream=raw, filetype="pdf")
    except Exception as exc:
        raise UnsupportedDocumentError(f"Could not open PDF: {exc}") from exc

    pages: list[Page] = []
    try:
        with _PYMUPDF_LOCK:
            total = doc.page_count
        limit = min(total, max_pages)
        for i in range(limit):
            with _PYMUPDF_LOCK:
                page = doc.load_page(i)
                # "text" mode = reading-order plain text; no layout boxes,
                # which is what a search or an LLM prompt wants.
                native = page.get_text("text") or ""
                pix = page.get_pixmap(dpi=render_dpi)
                # pix.samples is a flat RGB(A) buffer; copied out under the
                # lock with its shape, reshaped and converted outside it.
                samples = bytes(pix.samples)
                width, height, channels = pix.width, pix.height, pix.n
                del pix, page
            arr = np.frombuffer(samples, dtype=np.uint8).reshape(height, width, channels)
            if channels == 4:
                arr = arr[:, :, :3]
            bgr = arr[:, :, ::-1].copy()  # RGB → BGR
            pages.append(
                Page(
                    index=i,
                    image_bgr=bgr,
                    text=native,
                    text_source="native" if native.strip() else "none",
                    width=width,
                    height=height,
                )
            )
        truncated = max(0, total - limit)
    finally:
        with _PYMUPDF_LOCK:
            doc.close()
    return pages, truncated


def _svg_zoom(width_pt: float, height_pt: float, render_dpi: int, max_pixels: int) -> float:
    """The zoom an SVG page renders at: ``render_dpi / 72``, scaled down when
    that would exceed ``max_pixels``.

    An SVG's page rectangle is in points (MuPDF reads ``width`` / ``height``,
    falling back to the ``viewBox``, then to 612×792), so the render at the
    requested DPI is ``w·h·zoom²`` pixels. Past the cap the zoom becomes
    ``√(max_pixels / (w·h))`` — the largest render that fits — rather than
    an error: a big-but-legitimate drawing still loads, just less sharply.
    """
    zoom = render_dpi / 72.0
    area = width_pt * height_pt
    if area > 0 and area * zoom * zoom > max_pixels:
        zoom = math.sqrt(max_pixels / area)
    return zoom


def _ref_warnings(raw: bytes) -> list[str]:
    """One "not rendered" line per external reference, capped."""
    refs = find_external_refs(raw)
    lines = []
    for ref in refs[:_MAX_REF_WARNINGS]:
        what = "image" if ref.tag in ("image", "feImage") else f"<{ref.tag}> reference"
        url = ref.url if len(ref.url) <= _MAX_WARNING_URL else ref.url[:_MAX_WARNING_URL] + "…"
        lines.append(f"external {what} not rendered: {url}")
    if len(refs) > _MAX_REF_WARNINGS:
        lines.append(f"+{len(refs) - _MAX_REF_WARNINGS} more external reference(s) not rendered")
    return lines


def _load_svg(raw: bytes, render_dpi: int, max_pixels: int) -> tuple[list[Page], list[str]]:
    """Render an SVG as one page — image plus native text — like a 1-page PDF.

    MuPDF opens the SVG as a one-page document: ``get_text`` reads its
    ``<text>`` elements (so a text search hits words, not markup) and
    ``get_pixmap`` draws it on white. The zoom is ``_svg_zoom``'s — the
    requested DPI unless that breaks ``max_pixels``.

    External references are never drawn (MuPDF fetches nothing); each one is
    reported as a warning so a caller can tell a deliberately blank region
    from a logo that lived on someone's CDN.

    Args:
        raw:        SVG bytes.
        render_dpi: Raster density, as for a PDF page.
        max_pixels: Ceiling on ``width × height`` of the render.

    Returns:
        ``(pages, warnings)`` — always exactly one page.

    Raises:
        UnsupportedDocumentError: MuPDF cannot parse it, it has no drawable
            area (a zero width or height), or the render fails (including
            MuPDF's own ``FzErrorLimit``).
    """
    import numpy as np
    import pymupdf

    try:
        with _PYMUPDF_LOCK:
            doc = pymupdf.open(stream=raw, filetype="svg")
    except Exception as exc:
        raise UnsupportedDocumentError(f"Could not open SVG: {exc}") from exc

    try:
        with _PYMUPDF_LOCK:
            page = doc.load_page(0)
            rect = page.rect
            if rect.width <= 0 or rect.height <= 0:
                raise UnsupportedDocumentError(
                    f"SVG has no drawable area ({rect.width:g}×{rect.height:g} pt) — "
                    "give the root <svg> a width/height or a viewBox."
                )
            zoom = _svg_zoom(rect.width, rect.height, render_dpi, max_pixels)
            try:
                native = page.get_text("text") or ""
                pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
                samples = bytes(pix.samples)
                width, height, channels = pix.width, pix.height, pix.n
                del pix
            except Exception as exc:
                raise UnsupportedDocumentError(f"Could not render SVG: {exc}") from exc
            del page
    finally:
        with _PYMUPDF_LOCK:
            doc.close()

    if not width or not height:
        raise UnsupportedDocumentError("SVG rendered to an empty image.")
    arr = np.frombuffer(samples, dtype=np.uint8).reshape(height, width, channels)
    if channels == 4:
        arr = arr[:, :, :3]
    bgr = arr[:, :, ::-1].copy()  # RGB → BGR
    page_out = Page(
        index=0,
        image_bgr=bgr,
        text=native,
        text_source="native" if native.strip() else "none",
        width=width,
        height=height,
    )
    return [page_out], _ref_warnings(raw)


def _load_txt(raw: bytes) -> list[Page]:
    """One page carrying the decoded text; .txt has no rendered surface."""
    # utf-8-sig drops a BOM; errors="replace" keeps a mostly-valid file usable
    # rather than failing the whole request over one bad byte.
    text = raw.decode("utf-8-sig", errors="replace")
    return [Page(index=0, image_bgr=None, text=text, text_source="native")]


def _load_docx(raw: bytes) -> list[Page]:
    """Extract paragraphs and table cells from a .docx in document order.

    python-docx exposes ``document.paragraphs`` and ``document.tables`` as two
    flat lists, which loses their interleaving — a table between two headings
    would be appended at the end. Walking the body XML instead keeps reading
    order, which matters for both text search snippets and LLM context.

    Table rows are flattened to ``cell | cell | cell`` so a row reads as one
    line (``"System Size | 8.4 kW"`` stays searchable as a unit).

    Args:
        raw: .docx (OOXML ZIP) bytes.

    Returns:
        A single-page list; .docx has no page concept before layout.

    Raises:
        UnsupportedDocumentError: If python-docx cannot open the package.
    """
    import io

    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as exc:
        raise UnsupportedDocumentError(f"Could not open .docx: {exc}") from exc

    lines: list[str] = []
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, document).text.strip()
            if text:
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                # Skip fully empty rows (layout tables are full of them).
                if any(cells):
                    lines.append(" | ".join(cells))

    return [Page(index=0, image_bgr=None, text="\n".join(lines), text_source="native")]


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


def load_document(
    raw: bytes,
    filename: Optional[str] = None,
    content_type: Optional[str] = None,
    *,
    kind: Optional[DocumentKind] = None,
    max_pages: int = DEFAULT_MAX_PAGES,
    render_dpi: int = DEFAULT_RENDER_DPI,
    image_decoder: Optional[ImageDecoder] = None,
    keep_source: bool = False,
    max_render_pixels: int = DEFAULT_MAX_RENDER_PIXELS,
) -> Document:
    """Detect what ``raw`` is and load it into a ``Document``.

    Args:
        raw:           Complete file bytes.
        filename:      Original filename (recorded on the Document; never used
                       for detection).
        content_type:  Declared MIME type (error messages only).
        kind:          Skip detection and load as this kind. Only pass a value
                       produced by ``detect_kind`` on the same bytes.
        max_pages:     Page cap for PDFs; extra pages are counted in
                       ``Document.truncated_pages``.
        render_dpi:    Raster density for PDF page renders (and SVG renders,
                       up to ``max_render_pixels``).
        image_decoder: ``bytes → BGR ndarray`` used for the image kind.
                       Defaults to ``default_image_decoder``.
        keep_source:   Retain ``raw`` on the Document as ``source_bytes``.
                       Only needed when the caller will later ask
                       ``pdf_text_regions`` where a phrase sits — it has to
                       re-open the file. Costs the document's size in memory
                       for the life of the job, so leave it off otherwise.
        max_render_pixels:
                       Ceiling on one SVG render's ``width × height``; past it
                       the zoom shrinks to fit (``_svg_zoom``). PDFs ignore it.

    Returns:
        A ``Document`` with at least one page. An SVG's ``warnings`` list its
        external references that were not drawn.

    Raises:
        UnsupportedDocumentError: Unrecognised bytes, a legacy .doc, or a
            per-kind parse failure.
    """
    resolved_kind = kind or detect_kind(raw, filename=filename, content_type=content_type)
    decoder = image_decoder or default_image_decoder

    truncated = 0
    warnings: list[str] = []
    if resolved_kind == "image":
        pages = _load_image(raw, decoder)
    elif resolved_kind == "pdf":
        pages, truncated = _load_pdf(raw, max_pages=max_pages, render_dpi=render_dpi)
    elif resolved_kind == "svg":
        pages, warnings = _load_svg(raw, render_dpi=render_dpi, max_pixels=max_render_pixels)
    elif resolved_kind == "txt":
        pages = _load_txt(raw)
    elif resolved_kind == "docx":
        pages = _load_docx(raw)
    else:  # defensive: detect_kind only returns the five kinds above
        raise UnsupportedDocumentError(f"Unhandled document kind: {resolved_kind!r}")

    if not pages:
        raise UnsupportedDocumentError(
            f"{resolved_kind} document contained no pages — nothing to assess."
        )

    return Document(
        kind=resolved_kind,
        filename=filename or "",
        size_bytes=len(raw),
        content_type=guess_content_type(resolved_kind, raw),
        pages=pages,
        truncated_pages=truncated,
        source_bytes=raw if keep_source else None,
        warnings=warnings,
    )


def pdf_page_count(raw: bytes) -> int:
    """How many pages a PDF has, WITHOUT rendering or extracting any of them.

    Opening a PDF with PyMuPDF parses the cross-reference table, not the page
    content, so this is milliseconds even for a long document — cheap enough
    for a request handler that has to refuse a multi-page upload before a
    worker spends a render on it.

    Raises:
        UnsupportedDocumentError: PyMuPDF cannot open the stream.
    """
    import pymupdf

    with _PYMUPDF_LOCK:
        try:
            doc = pymupdf.open(stream=raw, filetype="pdf")
        except Exception as exc:
            raise UnsupportedDocumentError(f"Could not open PDF: {exc}") from exc
        try:
            return int(doc.page_count)
        finally:
            doc.close()


def pdf_page_size(
    raw: bytes, index: int = 0, *, render_dpi: int = DEFAULT_RENDER_DPI
) -> tuple[int, int]:
    """The ``(width, height)`` in pixels page ``index`` renders to at ``render_dpi``.

    The size ``load_document`` would give that page, WITHOUT rendering it:
    the page rectangle is transformed by the same zoom matrix
    ``get_pixmap(dpi=…)`` builds and its integer rectangle read off. Lets a
    request handler check caller-supplied pixel coordinates against a PDF
    page it has not rasterised yet.

    Raises:
        UnsupportedDocumentError: PyMuPDF cannot open the stream.
        IndexError: ``index`` is not a page of this PDF.
    """
    import pymupdf

    with _PYMUPDF_LOCK:
        try:
            doc = pymupdf.open(stream=raw, filetype="pdf")
        except Exception as exc:
            raise UnsupportedDocumentError(f"Could not open PDF: {exc}") from exc
        try:
            if not 0 <= index < doc.page_count:
                raise IndexError(f"page {index} is not in a {doc.page_count}-page PDF")
            page = doc.load_page(index)
            zoom = render_dpi / 72.0
            rect = page.rect * pymupdf.Matrix(zoom, zoom)
            irect = rect.irect
            return int(irect.width), int(irect.height)
        finally:
            doc.close()


def svg_page_size(
    raw: bytes,
    *,
    render_dpi: int = DEFAULT_RENDER_DPI,
    max_pixels: int = DEFAULT_MAX_RENDER_PIXELS,
) -> tuple[int, int]:
    """The ``(width, height)`` in pixels an SVG renders to — without rendering.

    ``pdf_page_size`` for the svg kind: the page rectangle MuPDF derives from
    ``width`` / ``height`` / ``viewBox``, transformed by the SAME zoom
    ``_load_svg`` uses (``_svg_zoom``, so the pixel cap applies identically)
    and its integer rectangle read off. A reference request needs it to check
    caller-supplied pixel regions against a page nobody has rasterised yet.

    Raises:
        UnsupportedDocumentError: MuPDF cannot open it, or it has no drawable
            area.
    """
    import pymupdf

    with _PYMUPDF_LOCK:
        try:
            doc = pymupdf.open(stream=raw, filetype="svg")
        except Exception as exc:
            raise UnsupportedDocumentError(f"Could not open SVG: {exc}") from exc
        try:
            rect = doc.load_page(0).rect
            if rect.width <= 0 or rect.height <= 0:
                raise UnsupportedDocumentError(
                    f"SVG has no drawable area ({rect.width:g}×{rect.height:g} pt)."
                )
            zoom = _svg_zoom(rect.width, rect.height, render_dpi, max_pixels)
            irect = (rect * pymupdf.Matrix(zoom, zoom)).irect
            return int(irect.width), int(irect.height)
        finally:
            doc.close()


# ---------------------------------------------------------------------------
# Native-PDF geometry for a text hit
# ---------------------------------------------------------------------------


def _normalise_token(token: str) -> str:
    """Lowercase and strip everything that is not a letter or digit.

    Word-span reconstruction has to compare the matcher's view of the text
    (``page.get_text("text")``, which carries punctuation and spacing) with
    PyMuPDF's word list (which splits on whitespace). Comparing on
    alphanumerics only is what makes ``"$4,850.00"`` line up with the word
    box PyMuPDF reports for it.
    """
    return "".join(ch for ch in token.lower() if ch.isalnum())


def _rect_for_snippet(
    words: list, snippet: str, start_index: int = 0
) -> tuple[Optional[tuple[float, float, float, float]], int]:
    """Find the consecutive word run whose text matches ``snippet``.

    Args:
        words:       PyMuPDF ``page.get_text("words")`` output —
                     ``(x0, y0, x1, y1, word, block, line, word_no)``.
        snippet:     The matched substring, as the text matcher saw it.
        start_index: Where to begin scanning, so repeated matches of the same
                     phrase return successive occurrences rather than the
                     first one over and over.

    Returns:
        ``(rect_or_None, next_start_index)``. The rect is the union of the
        run's word boxes, in PDF points.

    This is approximate by construction: a phrase that PyMuPDF splits
    differently from the text layer (hyphenation, a ligature, a column break
    mid-phrase) will not line up, and the function returns None rather than a
    plausible-looking wrong rectangle.
    """
    tokens = [t for t in (_normalise_token(w) for w in snippet.split()) if t]
    if not tokens:
        return None, start_index
    for i in range(start_index, len(words) - len(tokens) + 1):
        if all(
            _normalise_token(words[i + j][4]) == tokens[j] for j in range(len(tokens))
        ):
            run = words[i : i + len(tokens)]
            return (
                (
                    min(w[0] for w in run),
                    min(w[1] for w in run),
                    max(w[2] for w in run),
                    max(w[3] for w in run),
                ),
                i + len(tokens),
            )
    return None, start_index


def pdf_text_regions(
    pdf_bytes: bytes,
    page: Page,
    hits: list,
    *,
    pattern: str,
    mode: str = "contains",
    label: Optional[str] = None,
    max_regions: int = 200,
    filetype: str = "pdf",
) -> list:
    """Where a text hit sits on a native PDF page, in page-image pixels.

    Two strategies, picked by match mode — the split the regions plan calls
    for:

      ``contains`` / ``exact``  PyMuPDF's own ``page.search_for(pattern)``.
                                Fast and exact for a literal phrase; it is
                                case-insensitive and ignores word boundaries,
                                so an ``exact`` criterion's rectangles can be
                                a superset of its (word-anchored) matches.
      ``regex`` / ``fuzzy``     word-span reconstruction: each hit's matched
                                text is lined up against ``get_text("words")``
                                and the covering word boxes are unioned.
                                Approximate — see ``_rect_for_snippet``.

    Coordinates are converted from PDF points to the page's rendered pixels
    (the space every Region lives in), and the original rectangle is kept in
    ``attrs["pdf_rect"]`` so PDF tooling can use it directly.

    **Rotated pages.** ``search_for`` and ``get_text("words")`` answer in the
    UNROTATED page space, while the render — and ``page.rect`` — are rotated
    by the page's ``/Rotate``. Each rectangle is mapped through
    ``rotation_matrix`` before scaling; without it a ``/Rotate 270`` plan-set
    sheet put every box on the wrong axis. ``pdf_rect`` stays unrotated,
    because that is the space PDF tooling (annotations, ``search_for``) uses.

    **SVG.** Pass ``filetype="svg"`` and the same machinery answers for an
    SVG's ``<text>``: MuPDF opens it as a one-page document in points. The
    point → pixel scale is read off ``page.width / rect.width`` rather than
    assumed from a DPI, so it stays right when ``_load_svg`` capped the zoom.

    Args:
        pdf_bytes:   The original PDF (``Document.source_bytes``) — or SVG,
                     with ``filetype="svg"``.
        page:        The already-loaded page — supplies the index and the
                     rendered pixel size.
        hits:        ``TextHit`` objects for this page from
                     ``match_text(..., locate=True)``.
        pattern:     The literal searched for (contains/exact modes).
        mode:        contains | exact | regex | fuzzy.
        label:       Region label; defaults to ``pattern``.
        max_regions: Cap on regions returned.
        filetype:    "pdf" (default) | "svg" — the ``Document.kind`` the
                     bytes are.

    Returns:
        ``box`` regions with ``source="pdf-text"`` and ``score=1.0``. Empty
        when PyMuPDF is unavailable, the page is out of range, or nothing
        lined up — never an exception, because a missing overlay must not
        fail the criterion that produced it.
    """
    from ..vision.model import Region

    if not pdf_bytes or not page.width or not page.height:
        return []

    try:
        import pymupdf
    except ImportError:  # pragma: no cover - depends on env
        return []

    region_label = label or pattern
    regions: list = []
    with _PYMUPDF_LOCK:
        try:
            doc = pymupdf.open(stream=pdf_bytes, filetype=filetype)
        except Exception:
            return []
        try:
            if page.index >= doc.page_count:
                return []
            pdf_page = doc.load_page(page.index)
            rect = pdf_page.rect            # rotated — the render's frame
            sx = page.width / rect.width if rect.width else 1.0
            sy = page.height / rect.height if rect.height else 1.0
            to_rendered = pdf_page.rotation_matrix  # identity when /Rotate is 0

            found: list[tuple[tuple[float, float, float, float], str, float]] = []
            if mode in ("contains", "exact") and pattern:
                for quad in pdf_page.search_for(pattern)[:max_regions]:
                    found.append(((quad.x0, quad.y0, quad.x1, quad.y1), pattern, 1.0))
            else:
                words = pdf_page.get_text("words")
                cursor = 0
                for hit in hits[:max_regions]:
                    box, cursor = _rect_for_snippet(words, getattr(hit, "text", ""), cursor)
                    if box:
                        found.append((box, getattr(hit, "text", ""), getattr(hit, "ratio", 1.0)))

            for (x0, y0, x1, y1), text, ratio in found:
                shown = pymupdf.Rect(x0, y0, x1, y1) * to_rendered
                regions.append(
                    Region(
                        page=page.index,
                        kind="box",
                        points=[(shown.x0 * sx, shown.y0 * sy), (shown.x1 * sx, shown.y1 * sy)],
                        label=region_label,
                        score=1.0,
                        source="pdf-text",
                        attrs={
                            "text": text[:120],
                            "ratio": round(float(ratio), 4),
                            "pdf_rect": [
                                round(x0, 2), round(y0, 2), round(x1, 2), round(y1, 2)
                            ],
                        },
                    )
                )
        finally:
            doc.close()
    return regions
